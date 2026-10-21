"""Gera a versão PowerPoint (.pptx, editável) a partir do layout exportado por pptx_extrair.js.

Uso: python3 gerar_pptx.py layout.json saida.pptx
Texto vira caixa de texto, tabelas viram tabelas nativas, as notas de cada slide vão para as
notas do PowerPoint e os dois vídeos ficam embutidos (slides 34a e R1a).
"""
import json
import os
import re
import sys

from lxml import etree
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, MSO_AUTO_SIZE, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Emu, Pt

AQUI = os.path.dirname(os.path.abspath(__file__))
APRES = os.path.dirname(AQUI)
PX = 6350  # EMU por pixel (1920 px = 13,333 pol.)
FONTE = "Montserrat"
VIDEOS = {
    "s34v": ("video_capa.png", "Video_Theo_animacao.mp4"),
    "r1v": ("lara_capa.png", "Video_Lara_animacao.mp4"),
}
ALIGN = {"center": PP_ALIGN.CENTER, "right": PP_ALIGN.RIGHT, "end": PP_ALIGN.RIGHT}


def rgb(css):
    v = [int(float(x)) for x in re.findall(r"[\d.]+", css)[:3]]
    return RGBColor(*v)


def e(v):
    return Emu(int(round(v * PX)))


def notas(sid):
    s = open(os.path.join(AQUI, "slides", sid + ".html"), encoding="utf-8").read()
    m = re.search(r"<aside>(.*?)</aside>", s, re.S)
    if not m:
        return ""
    t = re.sub(r"<[^>]+>", "", m.group(1)).strip()
    return t.replace(" || ", "\n\n")


def add_runs(p, runs, base):
    for r in runs:
        if r["t"] == "\n":
            p.add_line_break()
            continue
        run = p.add_run()
        run.text = r["t"]
        f = run.font
        f.name = FONTE
        f.size = Pt(r.get("size", base["size"]) * 0.5)
        f.bold = r.get("b", False)
        f.italic = r.get("i", False)
        f.color.rgb = rgb(r.get("color", base["color"]))


def caixa(slide, it, extra_w=0.0):
    w = min(it["w"] * (1 + extra_w), 1920 - it["x"])
    tb = slide.shapes.add_textbox(e(it["x"]), e(it["y"]), e(w), e(it["h"]))
    tf = tb.text_frame
    tf.word_wrap = True
    tf.auto_size = MSO_AUTO_SIZE.NONE
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    tf.vertical_anchor = MSO_ANCHOR.TOP
    return tb, tf


def estilo_par(p, it):
    p.alignment = ALIGN.get(it["align"], PP_ALIGN.LEFT)
    p.line_spacing = round(it["lh"] / 1.22, 2)


def marcador(p, ordered, n):
    pPr = p._p.get_or_add_pPr()
    pPr.set("marL", str(int(44 * PX)))
    pPr.set("indent", str(int(-44 * PX)))
    if ordered:
        bu = etree.SubElement(pPr, qn("a:buAutoNum"))
        bu.set("type", "arabicPeriod")
        bu.set("startAt", "1")
    else:
        bu = etree.SubElement(pPr, qn("a:buChar"))
        bu.set("char", "•")


def tabela(slide, it):
    rows = it["rows"]
    ncol = max(len(r["cells"]) for r in rows)
    gs = slide.shapes.add_table(len(rows), ncol, e(it["x"]), e(it["y"]), e(it["w"]), e(it["h"]))
    tbl = gs.table
    tblPr = tbl._tbl.tblPr
    for k in ("firstRow", "bandRow"):
        tblPr.set(k, "0")
    style = tblPr.find(qn("a:tableStyleId"))
    if style is None:
        style = etree.SubElement(tblPr, qn("a:tableStyleId"))
    style.text = "{2D5ABB26-0587-4C30-8999-92F81FD0307C}"  # sem estilo, sem grade
    for j, c in enumerate(rows[0]["cells"]):
        tbl.columns[j].width = e(c["w"])
    for i, r in enumerate(rows):
        tbl.rows[i].height = e(max(c["h"] for c in r["cells"]))
        for j, c in enumerate(r["cells"]):
            cell = tbl.cell(i, j)
            bg = c["bg"] if "0, 0, 0, 0" not in c["bg"] else r["bg"]
            if "0, 0, 0, 0" not in bg:
                cell.fill.solid()
                cell.fill.fore_color.rgb = rgb(bg)
            else:
                cell.fill.background()
            cell.margin_left = cell.margin_right = e(c["pad"])
            cell.margin_top = cell.margin_bottom = e(c["pad"] * 0.6)
            cell.vertical_anchor = MSO_ANCHOR.MIDDLE
            tf = cell.text_frame
            tf.word_wrap = True
            p = tf.paragraphs[0]
            p.line_spacing = 1.0
            add_runs(p, c["runs"], c)


def main(layout, saida):
    data = json.load(open(layout, encoding="utf-8"))
    prs = Presentation()
    prs.slide_width, prs.slide_height = e(1920), e(1080)
    blank = prs.slide_layouts[6]
    for s in data:
        slide = prs.slides.add_slide(blank)
        bg = slide.background.fill
        bg.solid()
        bg.fore_color.rgb = rgb(s["bg"])
        for it in s["items"]:
            k = it["kind"]
            if k == "rect":
                shp = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE if it.get("radius") else MSO_SHAPE.RECTANGLE,
                                             e(it["x"]), e(it["y"]), e(it["w"]), e(it["h"]))
                shp.fill.solid()
                shp.fill.fore_color.rgb = rgb(it["bg"])
                shp.line.fill.background()
                shp.shadow.inherit = False
            elif k == "text":
                tb, tf = caixa(slide, it, 0.03)
                p = tf.paragraphs[0]
                estilo_par(p, it)
                add_runs(p, it["runs"], it)
            elif k == "list":
                tb, tf = caixa(slide, it, 0.03)
                for n, runs in enumerate(it["items"]):
                    p = tf.paragraphs[0] if n == 0 else tf.add_paragraph()
                    estilo_par(p, it)
                    marcador(p, it["ordered"], n)
                    if it.get("gap"):
                        p.space_before = Pt(it["gap"] * 0.5)
                    add_runs(p, runs, it)
            elif k == "table":
                tabela(slide, it)
            elif k == "img":
                capa, video = VIDEOS[s["id"]]
                mv = slide.shapes.add_movie(os.path.join(APRES, video), e(it["x"]), e(it["y"]), e(it["w"]), e(it["h"]),
                                            poster_frame_image=os.path.join(AQUI, capa), mime_type="video/mp4")
                mv.name = it.get("alt") or video
        n = notas(s["id"])
        if n:
            slide.notes_slide.notes_text_frame.text = n
    prs.core_properties.title = "Workshop ABA e TEA — LACC"
    prs.core_properties.author = "Betina Binatti"
    prs.save(saida)
    print(len(data), "slides ->", saida)


if __name__ == "__main__":
    main(sys.argv[1], sys.argv[2])
